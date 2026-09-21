"""
build_data.py — Tablero Competitividad Territorial Tolima & Huila (IDC 2026)
Lee la presentación IDC_2026_Tolima_-_Huila.pptx y genera data.json.
Uso: python3 build_data.py [ruta_pptx]
"""
import sys, re, json, unicodedata
from pptx import Presentation

SRC = sys.argv[1] if len(sys.argv) > 1 else "/mnt/user-data/uploads/IDC_2026_Tolima_-_Huila.pptx"
prs = Presentation(SRC)

def num(s):
    return float(s.replace('.', '').replace(',', '.')) if ',' in s else float(s)

def nfkc(s):
    return unicodedata.normalize('NFKC', s).strip()

def texts(slide):
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            out.append((sh.top // 9144, sh.left // 9144, sh.text_frame.text.strip()))
    return sorted(out)

def chart_of(slide):
    for sh in slide.shapes:
        if sh.has_chart:
            ch = sh.chart
            pl = ch.plots[0]
            return list(pl.categories), [list(se.values) for se in pl.series]
    return None, None

def kind(slide):
    t = texts(slide)
    return t[0][2] if t else ''

# ---- factores IDC (agrupación de los 13 pilares) ----
FACTORES = {
  'INS': 'Condiciones habilitantes', 'INF': 'Condiciones habilitantes',
  'TIC': 'Condiciones habilitantes', 'AMB': 'Condiciones habilitantes',
  'SAL': 'Capital humano', 'EDU': 'Capital humano', 'EDS': 'Capital humano',
  'NEG': 'Eficiencia de los mercados', 'LAB': 'Eficiencia de los mercados',
  'FIN': 'Eficiencia de los mercados', 'TAM': 'Eficiencia de los mercados',
  'SOF': 'Ecosistema innovador', 'INN': 'Ecosistema innovador',
}

pillars = {}
order = []

def get_p(code):
    if code not in pillars:
        pillars[code] = dict(code=code, subpillars=[], indicators=[], hist=None)
        order.append(code)
    return pillars[code]

cur = None
summary = []
for idx, slide in enumerate(prs.slides, 1):
    tx = texts(slide)
    if not tx: continue
    k = tx[0][2]
    if k.startswith('COMPARACIÓN GENERAL'):
        cats, ser = chart_of(slide)
        for i, c in enumerate(cats):
            m = re.match(r'(.+?) \[T(\d+)/H(\d+)\]', c)
            summary.append(dict(name=m.group(1), rankT=int(m.group(2)), rankH=int(m.group(3)),
                                T=ser[0][i], H=ser[1][i], nat=ser[2][i]))
        continue
    m = re.search(r'CAPÍTULO (\d+) · (\w+)', ' '.join(t[2] for t in tx[:3]))
    if not m: continue
    code = m.group(2)
    P = get_p(code)
    P['chapter'] = int(m.group(1))
    P['factor'] = FACTORES[code]

    if k.startswith('RESULTADO 2026'):
        P['name'] = tx[2][2]
        body = [t[2] for t in tx[3:]]
        # tarjetas superiores
        flat = [t for t in tx]
        def find_after(label, off=1):
            lab = [t for t in flat if t[2] == label][0]
            col = sorted([t for t in flat if abs(t[1] - lab[1]) < 12 and t[0] > lab[0]])
            return col[off - 1][2]
        P['T'] = dict(score=num(find_after('TOLIMA')), rank=int(find_after('TOLIMA', 2).lstrip('#')))
        P['H'] = dict(score=num(find_after('HUILA')), rank=int(find_after('HUILA', 2).lstrip('#')))
        P['nat'] = num(find_after('PROMEDIO NACIONAL'))
        P['top10'] = num(find_after('PROMEDIO TOP 10'))
        # subpilares
        i = 0
        while i < len(flat):
            if re.match(r'Tolima [\d,]+ · #\d+', flat[i][2]):
                nm = flat[i - 1][2]
                mt = re.match(r'Tolima ([\d,]+) · #(\d+)', flat[i][2])
                mh = re.match(r'Huila ([\d,]+) · #(\d+)', flat[i + 1][2])
                mp = re.match(r'Top 10 ([\d,]+)', flat[i + 2][2])
                P['subpillars'].append(dict(name=nm,
                    T=dict(score=num(mt.group(1)), rank=int(mt.group(2))),
                    H=dict(score=num(mh.group(1)), rank=int(mh.group(2))),
                    top10=num(mp.group(1))))
                i += 3
            else:
                i += 1
        for lab_txt, key in (('Mejor posición Tolima', 'bestT'), ('Mejor posición Huila', 'bestH')):
            lab = [t for t in flat if t[2] == lab_txt][0]
            cand = sorted([t for t in flat if abs(t[0] - lab[0]) < 15 and t[1] > lab[1] and re.search(r': #\d+$', t[2])], key=lambda t: t[1])
            mm = re.match(r'(.+): #(\d+)$', cand[0][2])
            P[key] = dict(name=mm.group(1), rank=int(mm.group(2)))
        for t in flat:
            mm = re.match(r'(\d+) indicadores', t[2])
            if mm: P['nInd'] = int(mm.group(1))

    elif k.startswith('HISTÓRICO'):
        cats, ser = chart_of(slide)
        P['hist'] = dict(years=[int(c) for c in cats], T=ser[0], H=ser[1], nat=ser[2], top10=ser[3])

    elif k.startswith('PUNTAJES'):
        cats, ser = chart_of(slide)
        for i, c in enumerate(cats):
            m2 = re.match(r'(.+?) \[T(\d+)/H(\d+)\] · (\d{4})', c)
            P['indicators'].append(dict(name=m2.group(1), year=int(m2.group(4)),
                T=dict(score=ser[0][i], rank=int(m2.group(2))),
                H=dict(score=ser[1][i], rank=int(m2.group(3))),
                nat=ser[2][i], top10=ser[3][i]))

    elif k.startswith('METODOLOGÍA'):
        # filas: se agrupan por posición vertical
        rows = []
        for top, left, t in tx:
            if top < 200 and left > 40 and top < 190: continue
            if left < 100 and top >= 200 and not t.startswith('El “año”'):
                rows.append(dict(name=t, top=top, calc='', fuente='', anio='', inv=False))
        for top, left, t in tx:
            if top < 200: continue
            if left < 100 or t.startswith('El “año”') or t.startswith('Fuente:'): continue
            # fila asociada = la última cuyo top <= top
            r = None
            for rr in rows:
                if rr['top'] <= top + 2: r = rr
            if r is None: continue
            if 300 <= left < 800: r['calc'] = t
            elif 800 <= left < 1120: r['fuente'] = t
            elif left >= 1120:
                if t.strip().upper() == 'INVERSO': r['inv'] = True
                else: r['anio'] = t
        P.setdefault('meta', []).extend(rows)

# ---- unir metodología con indicadores ----
def norm(s): return re.sub(r'\W+', '', nfkc(s).lower())
for code in order:
    P = pillars[code]
    meta = {norm(r['name']): r for r in P.get('meta', [])}
    for ind in P['indicators']:
        r = meta.get(norm(ind['name']))
        if r:
            ind['calc'] = nfkc(r['calc']).replace('\n', ' ')
            ind['fuente'] = nfkc(r['fuente']).replace('\n', ' ')
            ind['inverso'] = r['inv']
        else:
            ind['calc'] = ind['fuente'] = ''; ind['inverso'] = False
            print('SIN META:', code, ind['name'])
    P.pop('meta', None)
    if P['hist'] is None: print('SIN HIST', code)

# ---- nombres cortos para gráficos ----
SHORT = {
  'INS': 'Instituciones', 'INF': 'Infraestructura', 'TIC': 'Adopción de TIC',
  'AMB': 'Sostenibilidad ambiental', 'SAL': 'Salud', 'EDU': 'Educación básica y media',
  'EDS': 'Educación superior y FpT', 'NEG': 'Entorno para los negocios',
  'LAB': 'Mercado laboral', 'FIN': 'Sistema financiero', 'TAM': 'Tamaño del mercado',
  'SOF': 'Sofisticación y diversificación', 'INN': 'Innovación',
}
FULL = dict(SHORT)
FULL['EDS'] = 'Educación superior y formación para el trabajo'
for c in order:
    pillars[c]['short'] = SHORT[c]
    pillars[c]['full'] = FULL[c]

# ---- puntaje general = promedio simple de los 13 pilares (metodología IDC) ----
# Verificado: reproduce el promedio nacional publicado (5,07 en 2026; 5,01 en 2025)
# y el puntaje 2025 de Tolima (5,58).
years = pillars[order[0]]['hist']['years']
def mean_series(key):
    return [round(sum(pillars[c]['hist'][key][i] for c in order) / len(order), 4) for i in range(len(years))]
general = dict(years=years, T=mean_series('T'), H=mean_series('H'), nat=mean_series('nat'))

# ---- datos EXTERNOS a la presentación (fuentes públicas citadas) ----
# La presentación no incluye el puesto general ni el puntaje de Bogotá; se toman de:
#  - CPC & U. del Rosario (comunicado IDC 2026, jul-2026) y prensa regional (La Nación, Diario del Huila)
#  - Escalafón completo: resumen IDC 2026 de la Cámara de Comercio de Manizales por Caldas (CCMPC)
ESCALAFON = ["Bogotá, D.C.","Antioquia","Valle del Cauca","Risaralda","Santander","Atlántico","Cundinamarca",
  "Caldas","Quindío","Boyacá","Bolívar","Tolima","Norte de Santander","Huila","Archipiélago de San Andrés",
  "Meta","Casanare","Magdalena","Cauca","Cesar","Nariño","Córdoba","Caquetá","Sucre","Arauca","La Guajira",
  "Putumayo","Guaviare","Chocó","Amazonas","Guainía","Vichada","Vaupés"]
PUNTAJES_PUBLICADOS = {"Bogotá, D.C.":7.96,"Antioquia":6.90,"Valle del Cauca":6.43,"Risaralda":6.35,"Santander":6.32,
  "Caldas":6.02,"Quindío":5.99,"Chocó":3.70,"Amazonas":3.43,"Guainía":3.34,"Vichada":2.92,"Vaupés":2.81}
externo = dict(
  rank_general=dict(T={"2025":11,"2026":12}, H={"2025":15,"2026":14}),
  lider=dict(name="Bogotá, D.C.", score=7.96),
  promedio_nacional_publicado={"2025":5.01,"2026":5.07},
  escalafon=[dict(rank=i+1, name=n, score=PUNTAJES_PUBLICADOS.get(n)) for i,n in enumerate(ESCALAFON)],
)

out = dict(
    edicion=2026,
    pillars=[pillars[c] for c in order],
    general=general,
    externo=externo,
)
out['n_pillars'] = len(order)
out['n_subpillars'] = sum(len(pillars[c]['subpillars']) for c in order)
out['n_indicators'] = sum(len(pillars[c]['indicators']) for c in order)
json.dump(out, open('data.json', 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
print('pilares', out['n_pillars'], 'subpilares', out['n_subpillars'], 'indicadores', out['n_indicators'])
print('general T', general['T'][-2:], 'H', general['H'][-2:], 'nat', general['nat'][-2:])
